import pytest
import os
import shutil
from pathlib import Path
from app.models.context import SkillExecutionContext, UserSessionContext
from app.skills.office.spreadsheet_skills import SpreadsheetEditSkill, SpreadsheetReadSkill
from app.services.file_service import file_service
import openpyxl

@pytest.fixture(scope="module")
def workspace():
    test_dir = Path("test_workspace_excel")
    test_dir.mkdir(exist_ok=True)
    orig_ws = file_service.workspace_root
    file_service.workspace_root = test_dir.resolve()
    
    # Create dummy excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws['A1'] = "Test"
    ws['A2'] = "=1+1" # Formula
    
    dummy_xlsx = test_dir / "dummy.xlsx"
    wb.save(dummy_xlsx)
        
    yield test_dir
    file_service.workspace_root = orig_ws
    shutil.rmtree(test_dir, ignore_errors=True)

@pytest.fixture
def context():
    return SkillExecutionContext(session=UserSessionContext(user_id="test", user_name="Test"))

@pytest.mark.asyncio
async def test_spreadsheet_edit(workspace, context):
    skill = SpreadsheetEditSkill()
    operations = [
        {"action": "set_cell", "sheet": "Sheet1", "cell": "B1", "value": "Edited"},
        {"action": "add_sheet", "new_sheet_name": "Sheet2"}
    ]
    res = await skill.execute(context, filepath="dummy.xlsx", operations=operations)
    
    # Verify
    wb = openpyxl.load_workbook(workspace / "dummy.xlsx")
    assert "Sheet2" in wb.sheetnames
    assert wb["Sheet1"]["B1"].value == "Edited"
    assert wb["Sheet1"]["A1"].value == "Test"
    
@pytest.mark.asyncio
async def test_spreadsheet_read_offset(workspace, context):
    skill = SpreadsheetReadSkill()
    res = await skill.execute(context, filepath="dummy.xlsx", offset=0, limit=1)
    
    assert res["workbook"] == "dummy.xlsx"
    assert len(res["sheets"][0]["rows"]) == 1
