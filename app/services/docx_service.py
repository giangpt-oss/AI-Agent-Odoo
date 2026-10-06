from docx import Document
import os
from app.core.exceptions import SkillValidationError
from app.services.file_service import file_service

class DocxService:
    def read_docx(self, filepath: str) -> dict:
        safe_path = file_service.get_safe_path(filepath)
        try:
            doc = Document(safe_path)
            text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            return {"status": "success", "text": text, "paragraph_count": len(doc.paragraphs)}
        except Exception as e:
            raise SkillValidationError(f"Error reading docx: {e}")

    def create_docx(self, filepath: str, title: str, content: str) -> str:
        safe_path = file_service.get_safe_path(filepath)
        doc = Document()
        doc.add_heading(title, 0)
        doc.add_paragraph(content)
        doc.save(safe_path)
        return safe_path

    def append_docx(self, filepath: str, append_content: str) -> str:
        safe_path = file_service.get_safe_path(filepath)
        try:
            doc = Document(safe_path)
            doc.add_paragraph(append_content)
            doc.save(safe_path)
            return safe_path
        except Exception as e:
            raise SkillValidationError(f"Error editing docx: {e}")

    def replace_text_docx(self, filepath: str, old_text: str, new_text: str) -> str:
        safe_path = file_service.get_safe_path(filepath)
        try:
            doc = Document(safe_path)
            for p in doc.paragraphs:
                if old_text in p.text:
                    p.text = p.text.replace(old_text, new_text)
            doc.save(safe_path)
            return safe_path
        except Exception as e:
            raise SkillValidationError(f"Error editing docx: {e}")

docx_service = DocxService()
