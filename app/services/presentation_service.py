from pptx import Presentation
from app.core.exceptions import SkillValidationError
from app.services.file_service import file_service

class PresentationService:
    def read_pptx(self, filepath: str) -> dict:
        safe_path = file_service.get_safe_path(filepath)
        try:
            prs = Presentation(safe_path)
            slides_data = []
            for i, slide in enumerate(prs.slides):
                text_runs = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text_runs.append(shape.text)
                
                # notes_slide might be None
                notes = ""
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text
                    
                slides_data.append({
                    "slide_number": i + 1,
                    "text": "\n".join(text_runs),
                    "notes": notes
                })
            return {"status": "success", "slides": slides_data}
        except Exception as e:
            raise SkillValidationError(f"Error reading pptx: {e}")

    def create_pptx(self, filepath: str, title: str, subtitle: str) -> str:
        safe_path = file_service.get_safe_path(filepath)
        try:
            prs = Presentation()
            title_slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(title_slide_layout)
            title_shape = slide.shapes.title
            subtitle_shape = slide.placeholders[1]
            
            title_shape.text = title
            subtitle_shape.text = subtitle
            
            prs.save(safe_path)
            return safe_path
        except Exception as e:
            raise SkillValidationError(f"Error creating pptx: {e}")

    def add_slide(self, filepath: str, title: str, content: str) -> str:
        safe_path = file_service.get_safe_path(filepath)
        try:
            prs = Presentation(safe_path)
            bullet_slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(bullet_slide_layout)
            shapes = slide.shapes
            
            title_shape = shapes.title
            body_shape = shapes.placeholders[1]
            
            title_shape.text = title
            tf = body_shape.text_frame
            tf.text = content
            
            prs.save(safe_path)
            return safe_path
        except Exception as e:
            raise SkillValidationError(f"Error adding slide: {e}")

presentation_service = PresentationService()
