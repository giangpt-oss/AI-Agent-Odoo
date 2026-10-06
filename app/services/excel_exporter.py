import os
import re
import time
from pathlib import Path
from typing import Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXPORTS_DIR = Path(__file__).resolve().parent.parent.parent / "exports"
EXPORTS_DIR.mkdir(parents=True, exist_ok=True)


class ExcelExporter:
    """Tạo file bảng tính Excel (.xlsx) chuẩn doanh nghiệp với phong cách chuyên nghiệp."""

    @staticmethod
    def create_excel_report(
        title: str,
        headers: list[str],
        rows: list[list[Any]],
        sheet_name: str = "Dữ liệu",
        company_name: str = "CÔNG TY HOPITA - TRỢ LÝ ĐIỀU HÀNH AI",
    ) -> str:
        """Tạo file Excel có định dạng tiêu đề, màu sắc, viền và căn chỉnh tự động.
        Trả về đường dẫn tuyệt đối của file .xlsx đã tạo.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name[:30]

        # 1. Palette màu & Phông chữ doanh nghiệp
        FONT_FAMILY = "Segoe UI"
        title_font = Font(name=FONT_FAMILY, size=14, bold=True, color="1F4E78")
        meta_font = Font(name=FONT_FAMILY, size=9, italic=True, color="595959")
        header_font = Font(name=FONT_FAMILY, size=11, bold=True, color="FFFFFF")
        cell_font = Font(name=FONT_FAMILY, size=10, color="000000")

        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")

        thin_side = Side(border_style="thin", color="D9D9D9")
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        num_cols = max(len(headers), 1)

        # 2. Dòng 1: Tiêu đề công ty
        ws.cell(row=1, column=1, value=company_name).font = Font(name=FONT_FAMILY, size=9, bold=True, color="7F7F7F")
        ws.row_dimensions[1].height = 18

        # 3. Dòng 2: Tiêu đề báo cáo
        clean_title = title.upper()
        ws.cell(row=2, column=1, value=clean_title).font = title_font
        ws.row_dimensions[2].height = 26

        # 4. Dòng 3: Thời gian trích xuất
        time_part = time.strftime("%d/%m/%Y %H:%M:%S")
        now_str = f"Thời gian trích xuất: {time_part}"
        ws.cell(row=3, column=1, value=now_str).font = meta_font
        ws.row_dimensions[3].height = 16

        # 5. Dòng 5: Tiêu đề các cột (Header Table)
        start_row = 5
        ws.row_dimensions[start_row].height = 24
        for col_idx, header_text in enumerate(headers, 1):
            c = ws.cell(row=start_row, column=col_idx, value=str(header_text).strip())
            c.font = header_font
            c.fill = header_fill
            c.alignment = align_center
            c.border = cell_border

        # 6. Dòng 6+: Ghi dữ liệu
        for row_idx, row_data in enumerate(rows, start=start_row + 1):
            ws.row_dimensions[row_idx].height = 20
            is_zebra = (row_idx % 2 == 0)

            for col_idx in range(1, num_cols + 1):
                val = row_data[col_idx - 1] if col_idx - 1 < len(row_data) else ""
                cell = ws.cell(row=row_idx, column=col_idx)

                # Format theo kiểu dữ liệu
                if isinstance(val, (int, float)):
                    cell.value = val
                    if isinstance(val, float) or val >= 1000:
                        cell.number_format = "#,##0"
                    cell.alignment = align_right
                else:
                    cell.value = "" if val is None else str(val)
                    cell.alignment = align_left

                cell.font = cell_font
                cell.border = cell_border
                if is_zebra:
                    cell.fill = zebra_fill

        # 7. Tự động căn chỉnh độ rộng cột (Auto-fit width)
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Bỏ qua dòng 1, 2, 3 khi tính độ rộng
                if cell.row < start_row:
                    continue
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

        # 8. Lưu file với tên ASCII an toàn tuyệt đối
        import unicodedata
        ascii_title = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode("ascii").lower()
        safe_name = re.sub(r"[^a-zA-Z0-9_\-]", "_", ascii_title.strip())
        safe_name = re.sub(r"_+", "_", safe_name).strip("_")[:40] or "bao_cao"
        timestamp = time.strftime("%Y%m%d_%H%M%S")
        filename = f"{safe_name}_{timestamp}.xlsx"
        filepath = EXPORTS_DIR / filename
        wb.save(filepath)

        return str(filepath)


excel_exporter = ExcelExporter()
